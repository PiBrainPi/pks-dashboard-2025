import openpyxl
import json

folder = "/home/claw_01_rasbpi5_1/bka_pks2025"
results = {}

# 1. Gesamtzahlen deutsch vs. nichtdeutsch
try:
    wb_de = openpyxl.load_workbook(f"{folder}/12-T40-TV-deutsch.xlsx", read_only=True)
    ws_de = wb_de['T40']
    de_total = None
    for row in ws_de.iter_rows(min_row=1, max_col=4, values_only=True):
        if row[1] and 'insgesamt' in str(row[1]).lower():
            de_total = row[3]
            break
    wb_de.close()

    wb_nd = openpyxl.load_workbook(f"{folder}/17-T50-TV-nichtdeutsch.xlsx", read_only=True)
    ws_nd = wb_nd['T50']
    nd_total = None
    for row in ws_nd.iter_rows(min_row=1, max_col=4, values_only=True):
        if row[1] and 'insgesamt' in str(row[1]).lower():
            nd_total = row[3]
            break
    wb_nd.close()

    de_int = int(de_total) if de_total else 0
    nd_int = int(nd_total) if nd_total else 0
    gesamt = de_int + nd_int
    results['1_gesamtzahlen'] = {
        "deutsche_tv_gesamt": de_int,
        "nichtdeutsche_tv_gesamt": nd_int,
        "gesamt_tv": gesamt,
        "anteil_deutsch_prozent": round(de_int / gesamt * 100, 2) if gesamt else None,
        "anteil_nichtdeutsch_prozent": round(nd_int / gesamt * 100, 2) if gesamt else None
    }
except Exception as e:
    results['1_gesamtzahlen'] = {"error": str(e)}

# 2. Top 10 Staatsangehörigkeiten (T62)
try:
    wb62 = openpyxl.load_workbook(f"{folder}/22-T62-TV-Staatsangehoerigkeiten.xlsx", read_only=True)
    ws62 = wb62['T62 BKA']

    headers = []
    for row in ws62.iter_rows(min_row=5, max_row=5, values_only=True):
        headers = list(row)
        break

    total_row = None
    for row in ws62.iter_rows(min_row=6, max_row=6, values_only=True):
        total_row = list(row)
        break

    nd_total_t62 = total_row[3] if len(total_row) > 3 else None

    laender = []
    for col_idx in range(4, len(total_row)):
        land_name = str(headers[col_idx]).strip() if col_idx < len(headers) else f"Col{col_idx}"
        wert = total_row[col_idx]
        if wert and isinstance(wert, (int, float)) and wert > 0:
            laender.append({"land": land_name, "tatverdaechtige": int(wert)})

    laender_sorted = sorted(laender, key=lambda x: x['tatverdaechtige'], reverse=True)[:10]
    results['2_top_staatsangehoerigkeiten'] = {
        "nichtdeutsche_tv_gesamt_t62": int(nd_total_t62) if nd_total_t62 else None,
        "top_10": laender_sorted
    }
    wb62.close()
except Exception as e:
    results['2_top_staatsangehoerigkeiten'] = {"error": str(e)}

# 3. Aufenthaltsanlass (T61)
try:
    wb61 = openpyxl.load_workbook(f"{folder}/23-T61-TV-nichtdeutsch-Aufenthaltsanlass.xlsx", read_only=True)
    ws61 = wb61['T61']

    aufenthalt_stats = {}
    for row in ws61.iter_rows(min_row=1, max_col=15, values_only=True):
        if row[1] and 'insgesamt' in str(row[1]).lower():
            # Spalte E=Anzahl TV, F=Anteil %, G=Kein Aufenthalt, H=unerlaubt
            aufenthalt_stats = {
                "nichtdeutsche_tv_gesamt": int(row[4]) if row[4] else None,
                "anteil_an_tv_gesamt_prozent": float(row[5]) if row[5] else None,
                "kein_aufenthalt_unerlaubt_anzahl": int(row[6]) if len(row)>6 and row[6] else None,
                "unerlaubt_aufenthalt_anzahl": int(row[7]) if len(row)>7 and row[7] else None,
            }
            break
    wb61.close()
    results['3_aufenthaltsanlass'] = aufenthalt_stats
except Exception as e:
    results['3_aufenthaltsanlass'] = {"error": str(e)}

# 4. Mehrfachtäter deutsch vs nichtdeutsch
try:
    wb_mf_de = openpyxl.load_workbook(f"{folder}/14-T40-Mehrfach-TV-deutsch.xlsx", read_only=True)
    ws_mf_de = wb_mf_de['T40 MFTV']
    mf_de_total = None
    for row in ws_mf_de.iter_rows(min_row=1, max_col=5, values_only=True):
        if row[1] and 'insgesamt' in str(row[1]).lower():
            mf_de_total = row[4]
            break
    wb_mf_de.close()

    wb_mf_nd = openpyxl.load_workbook(f"{folder}/19-T50-Mehrfach-TV-nichtdeutsch.xlsx", read_only=True)
    ws_mf_nd = wb_mf_nd['T50 MFTV']
    mf_nd_total = None
    for row in ws_mf_nd.iter_rows(min_row=1, max_col=5, values_only=True):
        if row[1] and 'insgesamt' in str(row[1]).lower():
            mf_nd_total = row[4]
            break
    wb_mf_nd.close()

    mf_de_int = int(mf_de_total) if mf_de_total else 0
    mf_nd_int = int(mf_nd_total) if mf_nd_total else 0
    results['4_mehrfachtäter'] = {
        "deutsche_mehrfachtv_gesamt": mf_de_int,
        "nichtdeutsche_mehrfachtv_gesamt": mf_nd_int,
        "verhaeltnis_de_zu_nd": round(mf_de_int / mf_nd_int, 2) if mf_nd_int else None
    }
except Exception as e:
    results['4_mehrfachtäter'] = {"error": str(e)}

print(json.dumps(results, indent=2, ensure_ascii=False))
